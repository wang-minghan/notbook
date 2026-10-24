from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pathlib import Path

ROOT=Path(r'C:/Users/32027/OneDrive/notebook/02_Areas/工作经历/鑫犇电子有限公司/06_公司公共/公司介绍PPT')
OUT=ROOT/'build'; OUT.mkdir(exist_ok=True)
C={'bg':'111820','blue':'16324A','cyan':'36C2C9','white':'F5F8FA','muted':'B5C2D0','light':'F4F7FA','ink':'1F2D3D','line':'D6E0E8','amber':'F6C453','teal':'1C9C8C','purple':'7967B8','red':'C95A5A'}
def rgb(x): return RGBColor.from_string(x)
def shape(sl,x,y,w,h,fill=None,line=None,rad=True):
    s=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=rgb(C[fill] if fill else ('bg' if fill is None else fill)); s.line.color.rgb=rgb(C[line] if line else (C[fill] if fill else C['line'])); return s
def text(sl,x,y,w,h,t,size=16,color='ink',bold=False,align=PP_ALIGN.LEFT):
    s=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=s.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(.08);tf.margin_right=Inches(.08);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0];p.text=t;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(C[color]);p.alignment=align;return s
def base(prs,dark=False):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(C['bg'] if dark else C['light']); return sl
def head(sl,k,t,sub='',dark=False):
    text(sl,.6,.28,12,.25,k,10,'cyan' if dark else 'teal',True);text(sl,.6,.62,12,.55,t,27,'white' if dark else 'ink',True);text(sl,.6,1.18,12,.3,sub,12,'muted' if dark else 'ink')
def foot(sl,s,dark=False): text(sl,.6,7.1,12,.18,s,8,'muted' if dark else 'ink')
def card(sl,x,y,w,h,title,body,color='blue',dark=False):
    shape(sl,x,y,w,h,color if dark else 'white',color if dark else 'line');text(sl,x+.12,y+.12,w-.24,.3,title,15,'white' if dark else color,True);text(sl,x+.12,y+.52,w-.24,h-.62,body,12,'white' if dark else 'ink')
def node(sl,x,y,w,h,t,color='blue',dark=False):
    shape(sl,x,y,w,h,color if dark else 'white',color);text(sl,x+.05,y+.04,w-.1,h-.08,t,14,'white' if dark else color,True,PP_ALIGN.CENTER)
def arrow(sl,x1,y1,x2,y2,color='cyan'):
    l=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2));l.line.color.rgb=rgb(C[color]);l.line.width=Pt(2);l.line.end_arrowhead=True

def build():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    hero=ROOT/'visual_assets'/'company_hero_v01.jpg'
    sl=base(prs,True);sl.shapes.add_picture(str(hero),Inches(6.55),0,width=Inches(6.78),height=Inches(7.5));text(sl,.65,.45,5.4,.3,'XINBEN ELECTRONICS',11,'cyan',True);text(sl,.65,1.0,5.5,.8,'鑫犇电子有限公司',30,'white',True);text(sl,.65,1.9,5.3,.5,'精密金属零件与冲压制造能力介绍',20,'white',True);text(sl,.65,3.0,5.2,1.0,'面向产品、工程、制造与品质协同的\n制造能力底稿',23,'white',True);text(sl,.65,6.42,5.4,.28,'AI 生成概念视觉，不代表真实工厂或设备',10,'amber');foot(sl,'对外介绍草案｜正式公司名称与定位待确认',True)
    sl=base(prs);head(sl,'02 / OVERVIEW','我们是谁','公司能力应从四条互相咬合的线理解，而不是单一产品或单一工序。')
    for x,y,t,b,c in [(.75,2.0,'产品','机壳、磁性/结构件、制动片及配套零件','blue'),(3.85,2.0,'工程','图纸版本、材料、模具与工程准备','purple'),(.75,4.25,'制造','冲压、清洗、生产与关键资源协同','teal'),(3.85,4.25,'品质','尺寸、外观、表面处理、单重与异常闭环','amber')]: card(sl,x,y,2.75,1.75,t,b,c)
    card(sl,7.2,2.0,5.25,4.0,'事实边界','正式对外全称与定位：待确认\n规模、产能、精度、良率、认证：不在本版宣称\n\n资料依据：公司事实底稿、产品信息总览、组织与职责矩阵','blue')
    foot(sl,'资料依据：OneDrive 公司资料｜未确认信息保持待确认')
    sl=base(prs);head(sl,'03 / PRODUCTS','产品与应用','产品类别按已验证资料归纳；公开范围与具体应用仍需确认。')
    for x,t,b,c in [(.75,'机壳与盖板','结构件、隔板、极片、垫片等','blue'),(4.25,'磁性/结构件','磁扣、磁抱、磁栅等类别','purple'),(7.75,'制动与配套','制动片及其他配套零件','teal')]: card(sl,x,2.0,3.0,2.55,t,b+'\n\n公开产品范围：待确认',c)
    shape(sl,.75,5.15,10.0,.7,'amber','amber');text(sl,.95,5.28,9.6,.38,'不展示敏感客户名称；具体型号、材料和单重按客户版本管理。',14,'ink',True);foot(sl,'资料依据：产品信息总览｜公开表述待确认')
    sl=base(prs);head(sl,'04 / MATERIALS','材料与版本管理','客户版本、图纸、材料、单重与检验共同支撑产品一致性。')
    xs=[.8,3.25,5.7,8.15,10.6];ts=['客户版本','受控图纸','材料规格','单重记录','检验记录']
    for i,t in enumerate(ts): node(sl,xs[i],2.55,1.75,1.0,t,'blue' if i<2 else 'teal');
    for i in range(4): arrow(sl,xs[i]+1.75,3.05,xs[i+1],3.05)
    card(sl,1.0,4.35,4.8,1.25,'控制重点','同型号不同客户版本可能对应不同材料和单重。','purple');card(sl,7.0,4.35,4.8,1.25,'对外边界','敏感客户信息不在本版展示；公开材料范围待确认。','amber');foot(sl,'资料依据：产品信息总览、材料型号资料')
    sl=base(prs);head(sl,'05 / ENGINEERING','模具与工程技术','模具、工程技术、生产与品质共同支撑制造交付。')
    for i,(t,c) in enumerate([('图纸评审','blue'),('工程准备','purple'),('模具维护','teal'),('生产验证','amber')]): node(sl,1.0+i*3.0,2.4,2.25,1.0,t,c); 
    for i in range(3): arrow(sl,3.25+i*3,2.9,4.0+i*3,2.9)
    card(sl,1.0,4.35,5.0,1.35,'能力证据','工程技术组、模具维修组及生产/品质协同为已知组织线索。','blue');card(sl,7.0,4.35,5.0,1.35,'待补证据','真实 CAD、模具和现场照片需归档确认；概念图不代表量产证明。','amber');foot(sl,'资料依据：生产部资料、CAD项目｜能力与设备边界待确认')
    sl=base(prs);head(sl,'06 / FLOW','订单到交付','订单从需求评审到交付，需要跨部门协同；责任与周期待确认。')
    lanes=[('客户/业务','需求与变更','blue'),('工程/模具','评审与准备','purple'),('生产','制造与记录','teal'),('品质','检验与放行','amber'),('仓储/交付','入库与交付','blue')]
    for i,(a,b,c) in enumerate(lanes): text(sl,.75,1.9+i*.82,1.55,.32,a,13,c,True);node(sl,2.55+i*2.05,1.82+i*.82,1.55,.52,b,c); 
    for i in range(4): arrow(sl,4.1+i*2.05,2.08+i*.82,4.6+i*2.05,2.08+i*.82)
    shape(sl,.75,6.1,11.7,.5,'amber','amber');text(sl,.95,6.18,11.3,.3,'流程节点、责任归属、合作周期：待正式流程文件确认',13,'ink',True);foot(sl,'资料依据：组织矩阵、流程资料｜不写未经确认的交期承诺')
    sl=base(prs);head(sl,'07 / QUALITY','质量控制','质量关注点贯穿图纸、材料、关键尺寸、外观、表面处理和异常闭环。')
    for i,(t,c) in enumerate([('图纸/版本','blue'),('材料/批次','purple'),('尺寸/单重','teal'),('外观/表面','amber')]): node(sl,.9+i*3.0,2.3,2.25,.85,t,c)
    for i in range(3): arrow(sl,3.15+i*3,2.72,3.9+i*3,2.72)
    arrow(sl,11.0,3.15,6.3,4.15,'cyan');text(sl,8.0,3.35,2.5,.3,'检查完成后进入放行判断',11,'teal',True)
    node(sl,5.15,4.15,2.3,.9,'放行？','teal');text(sl,7.65,4.32,.5,.25,'是',11,'teal',True);arrow(sl,7.45,4.6,9.2,4.6,'cyan');node(sl,9.2,4.15,2.3,.9,'交付记录','blue')
    text(sl,4.25,5.2,.5,.25,'否',11,'red',True);arrow(sl,5.15,5.05,3.25,5.55,'red');node(sl,1.0,5.25,2.3,.75,'隔离 / 评审 / 整改','red');arrow(sl,3.3,5.62,5.15,5.62,'red');node(sl,5.15,5.25,2.3,.75,'复检记录','purple');text(sl,7.8,5.23,3.7,.75,'复检通过 → 返回放行判断',12,'purple',True);foot(sl,'指标、阈值和正式检验表单：待确认')
    sl=base(prs);head(sl,'08 / TRACEABILITY','产品一致性与追溯','版本、批次、产品编码和检验记录形成可回溯链路。')
    for i,(t,c) in enumerate([('图纸版本','blue'),('材料批次','purple'),('产品编码','teal'),('检验记录','amber')]): node(sl,1.0+i*3.0,2.35,2.3,.9,t,c)
    for i in range(3): arrow(sl,3.3+i*3,2.8,4.0+i*3,2.8)
    card(sl,1.0,4.35,5.0,1.35,'记录对象','版本清单、材料/批次记录、产品编码、检验记录。','blue');card(sl,7.0,4.35,5.0,1.35,'待补样例','脱敏真实记录模板待补；不使用伪造样本。','amber');foot(sl,'资料依据：质量/运营模板｜字段与保存期限待确认')
    sl=base(prs);head(sl,'09 / FACTORY','现场能力','真实现场素材是展示制造能力的证据，本页暂保留实拍待补位。')
    for i,(t,c) in enumerate([('工厂实拍\n待补','blue'),('设备实拍\n待补','purple'),('模具实拍\n待补','teal'),('检验实拍\n待补','amber')]): node(sl,.75+i*3.05,2.15,2.35,2.0,t,c)
    shape(sl,1.0,4.75,10.8,.7,'amber','amber');text(sl,1.25,4.92,10.3,.32,'不以 AI 图替代真实工厂、设备、模具或检验照片；需取得授权后归档。',14,'ink',True);foot(sl,'现场照片、授权和来源：待补')
    sl=base(prs);head(sl,'10 / COOPERATION','合作方式','以需求评审、试制、量产、交付为主线组织合作。')
    for i,(t,c) in enumerate([('需求输入','blue'),('工程评审','purple'),('样件/试制','teal'),('量产交付','amber')]): node(sl,.9+i*3.0,2.55,2.3,1.0,t,c)
    for i in range(3): arrow(sl,3.2+i*3,3.05,3.9+i*3,3.05)
    card(sl,1.0,4.55,4.9,1.2,'客户提供','图纸/版本、材料要求、质量要求和交付需求。','blue');card(sl,7.0,4.55,4.9,1.2,'待确认承诺','合作周期、样件节奏、产能与交付承诺需正式确认。','amber');foot(sl,'资料依据：订单/供应商模板｜不写未确认周期')
    sl=base(prs);head(sl,'11 / ROADMAP','能力扩展方向','只展示已确认的建设事项，并明确当前状态。')
    for x,t,b,c in [(.8,'现状','已验证的产品、组织与质量关注点','blue'),(4.25,'建设中','有正式资料支持、正在补齐的证据','teal'),(7.7,'拟建立','规划事项，不能写成当前能力','amber')]: card(sl,x,2.05,3.0,2.6,t,b+'\n\n具体事项：待确认',c)
    shape(sl,.8,5.25,9.9,.55,'amber','amber');text(sl,1.0,5.36,9.5,.3,'每项规划必须标注状态、负责人和证据来源；不使用“行业领先”等无证据表述。',13,'ink',True);foot(sl,'规划资料与状态：待确认')
    sl=base(prs,True);head(sl,'12 / CONTACT','合作入口','正式公司信息确认后，再开放对外联系方式。',True)
    card(sl,1.0,2.0,5.2,2.8,'正式公司信息','公司全称：待确认\n对外简称：待确认\n地址：待确认\n电话 / 邮箱：待确认','blue',True);card(sl,7.0,2.0,4.5,2.8,'二维码与入口','二维码：待补\n联系人：待确认\n\n不生成虚假联系方式。','teal',True);foot(sl,'对外介绍草案｜所有联系信息逐项核验后发布',True)
    prs.save(OUT/'xinben_company_intro_v02.pptx');print(OUT/'xinben_company_intro_v02.pptx',len(prs.slides))
if __name__=='__main__': build()
