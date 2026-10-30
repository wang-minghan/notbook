from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT=Path(r'C:\Users\32027\.hermes\workspace\鑫犇电子岗位操作与知识手册')
OUT=ROOT/'鑫犇电子岗位操作与知识手册_第二版.pptx'
XLS=ROOT/'鑫犇电子_现场速查表_第二版.xlsx'
NAVY=RGBColor(20,43,65); BLUE=RGBColor(42,105,145); TEAL=RGBColor(41,137,131); ORANGE=RGBColor(232,140,50); INK=RGBColor(35,44,52); MUTED=RGBColor(95,107,116); PALE=RGBColor(238,245,247); WHITE=RGBColor(255,255,255); RED=RGBColor(177,64,57)

# compact field workbook
sheets={
'01任务确认':['日期','工单号','客户','产品名称','型号','图号/版本','计划数量','交期','材料牌号','厚度','宽度','模具编号','设备编号','上次异常','四项确认结果','确认人'],
'02领料核对':['日期','工单号','产品编码','材料牌号','厚度','宽度','批次','领料重量','余料重量','标签状态','结果','核对人'],
'03首件检查':['日期','工单号','产品编码','检查项目','图纸标准','实测值','单位','外观结果','结果','检查人','确认人'],
'04生产巡检':['日期','工单号','产品编码','时间','累计数量','检查项目','标准','实测值','异常情况','处理结果','检查人'],
'05异常处理':['日期时间','工单号','产品编码','模具编号','工位号','异常现象','发现数量','隔离范围','先行处理','通知部门','恢复条件','结果','记录人'],
'06易混材料':['客户','产品名称','型号','材料牌号','厚度','宽度','易混对象','区别','核对依据','状态'],
'07镶件台账':['零件ID','模具编号','工位号','零件名称','零件类型','材料','关键尺寸','方向特征','图纸/版本','备品数量','位置','状态'],
'08试模记录':['日期','试模编号','模具编号','产品编码','工位号','问题现象','调整内容','调整前','调整后','样件结果','后续处理','记录人'],
'09修机报告':['日期','报告编号','模具编号','产品编码','工位号','故障现象','原因','处理方式','更换零件ID','试模编号','试模结果','后续提升','维修人'],
'10模具备品':['零件ID','模具编号','工位号','零件名称','最低库存','当前库存','存放位置','领用日期','补充状态','备注']}
wb=Workbook(); wb.remove(wb.active)
navy='142B41'; pale='EAF2F4'; orange='FCE4C3'; thin=Side(style='thin',color='B8C4C8')
for name,headers in sheets.items():
    ws=wb.create_sheet(name); ws.freeze_panes='A2'; ws.sheet_view.showGridLines=False
    for i,h in enumerate(headers,1):
        c=ws.cell(1,i,h); c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=navy); c.alignment=Alignment(horizontal='center',wrap_text=True); c.border=Border(bottom=thin)
        ws.column_dimensions[c.column_letter].width=min(24,max(12,len(h)*2+4))
    ws.auto_filter.ref=f'A1:{ws.cell(1,len(headers)).coordinate}'
    for r in range(2,102):
        for c in range(1,len(headers)+1): ws.cell(r,c).alignment=Alignment(vertical='top',wrap_text=True)
    if name=='07镶件台账':
        rows=[['EX-P-001','EX-M-001','OP10','Φ2.0导正孔冲头','冲头','SKH51','直径/总长/挂台待图纸确认','挂台朝上（示例）','EX-DWG-P001/A',2,'示例柜A1','示例/待确认'],['EX-D-001','EX-M-001','OP30','C型外圆冲头','冲头','DC53','外圆/开口槽/刃口待确认','开口方向按排样图','EX-DWG-D001/A',1,'示例柜A2','示例/待确认']]
        for row in rows: ws.append(row)
    if name=='08试模记录':
        ws.append(['2026-09-18','EX-TR-001','EX-M-001','EX-LDW-1027-上机壳','OP10','导正孔位置偏移（虚拟示例）','调整送料步距并复查导正销','偏移0.08','偏移0.01','示例：合格','连续样件复核','示例'])
        ws.append(['2026-09-18','EX-TR-002','EX-M-001','EX-LDW-1027-上机壳','OP50','折弯高度偏高（虚拟示例）','调整成型镶件垫片','高0.06','高0.01','示例：合格','记录垫片组合','示例'])
    if name=='09修机报告':
        ws.append(['2026-09-18','EX-RP-001','EX-M-001','EX-LDW-1027-上机壳','OP20','小孔毛刺突然增大（虚拟示例）','刃口局部崩角（示例）','更换备用冲头并试模','EX-P-002','EX-TR-003','示例：合格','检查冲头寿命记录','示例'])
        ws.append(['2026-09-18','EX-RP-002','EX-M-001','EX-LDW-1027-上机壳','OP60','整形后尺寸波动（虚拟示例）','成型镶件松动（示例）','重新定位紧固并复核','', 'EX-TR-004','示例：合格','补充紧固检查点','示例'])
    if ws.max_row>1:
        for r in range(2,ws.max_row+1):
            for c in range(1,len(headers)+1): ws.cell(r,c).fill=PatternFill('solid',fgColor='FFF2CC')
    ref=f'A1:{ws.cell(101,len(headers)).coordinate}'; tab=Table(displayName='T'+str(len(wb.worksheets)).zfill(2),ref=ref); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True); ws.add_table(tab)
wb.save(XLS)

# PPT helpers
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
def rect(slide,x,y,w,h,fill,outline=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=fill; shape.line.color.rgb=outline or fill
    return shape
def text(slide,txt,x,y,w,h,size=18,color=INK,bold=False,font='Microsoft YaHei',align=PP_ALIGN.LEFT,margin=0.08):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(margin); tf.margin_right=Inches(margin); tf.margin_top=Inches(margin); tf.margin_bottom=Inches(margin); tf.vertical_anchor=MSO_ANCHOR.TOP
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=txt; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box
def add_header(slide,kicker,title,num):
    text(slide,kicker.upper(),0.65,0.35,4,0.28,10,BLUE,True); text(slide,title,0.65,0.7,11.5,0.55,27,NAVY,True); text(slide,f'{num:02d}',12.35,0.4,0.45,0.3,12,MUTED,True,align=PP_ALIGN.RIGHT)
def bullets(slide,items,x,y,w,h,size=16,color=INK):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=box.text_frame; tf.clear(); tf.word_wrap=True; tf.margin_left=Inches(0.05); tf.margin_right=Inches(0.05)
    for i,item in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=item; p.level=0; p.font.name='Microsoft YaHei'; p.font.size=Pt(size); p.font.color.rgb=color; p.space_after=Pt(10)
    return box
def footer(slide,txt='鑫犇电子有限公司｜岗位操作与知识手册 第二版'):
    text(slide,txt,0.65,7.12,10,0.18,8,MUTED)
def slide_title(title,sub,idx):
    s=prs.slides.add_slide(blank); rect(s,0,0,13.333,7.5,WHITE); rect(s,0,0,0.18,7.5,TEAL); add_header(s,'鑫犇电子｜岗位操作与知识手册',title,idx); text(s,sub,0.7,1.38,11.7,0.42,13,MUTED); footer(s); return s

def cover():
 s=prs.slides.add_slide(blank); rect(s,0,0,13.333,7.5,NAVY); rect(s,0,0,4.2,7.5,BLUE); rect(s,0,6.86,13.333,0.64,ORANGE); text(s,'XINBEN / WORK MANUAL',0.75,0.65,5,0.35,14,WHITE,True); text(s,'鑫犇电子有限公司\n岗位操作与知识手册',0.75,1.55,8.8,1.5,34,WHITE,True); text(s,'给新员工看的公司工作地图\n认识部门｜看懂流程｜知道遇到问题找谁',0.82,3.55,7.5,0.85,18,RGBColor(220,232,236)); rect(s,8.9,1.55,3.25,3.25,RGBColor(27,67,91),RGBColor(75,154,164)); text(s,'客户订单',9.25,2.0,2.5,0.4,20,WHITE,True,align=PP_ALIGN.CENTER); text(s,'↓',10.2,2.48,1,0.5,28,ORANGE,True,align=PP_ALIGN.CENTER); text(s,'生产｜模具\n↓\n首件｜量产\n↓\n完工交付',9.35,3.05,2.35,1.3,18,WHITE,True,align=PP_ALIGN.CENTER); text(s,'第二版｜实用版',0.8,6.98,4,0.25,11,NAVY,True); return s
cover()
slides=[
('使用说明','这不是大厂制度汇编，而是新员工的工作地图',['先看第一册：知道公司做什么、三个部门怎么配合。','再看自己所在部门：知道本岗位负责什么、遇到问题找谁。','附件是现场速查表，现场需要时直接拿来填。']),
('公司目前的三个部门','三部门各管一段，互相接上才完成一张订单',['生产部：把正确的产品做出来，并负责现场检查与记录。','模具部：让模具可用、好用，处理模具造成的异常。','行政管理部：负责订单和资料流转、人员及日常综合事务。']),
('一张客户订单怎样完成','一条线看懂公司工作从哪里开始、到哪里结束',['1 收到订单｜行政管理部／公司负责人（实际接单人待确认）','2 确认能否生产｜生产部负责，模具部配合','3 安排生产｜生产部','4 准备材料｜生产部','5 准备模具｜模具部','6 首件确认 → 7 批量生产 → 8 异常处理 → 9 完工交付']),
('订单流程中，问题在哪一段就找谁','不绕圈，不让新员工自己猜',['产品、材料、数量、排产：先找生产部。','模具、工位、试模、修模：先找模具部。','订单资料、文件、考勤、用品：先找行政管理部。','跨部门或说不清：先保留现场和记录，找公司负责人（具体人待确认）。']),
('生产部负责什么','生产部的结果不只是“机器开了”',['产品做对：客户、型号、图纸和材料对应。','过程看住：首件、关键尺寸、外观和数量。','状态分清：合格品、不良品、待判品不混放。','异常说清：什么时候发现、影响多少、需要谁处理。']),
('生产部：开工前四确认','四项对不上，不凭经验开机',['工单对产品：客户、产品名称、型号、数量。','图纸对版本：图号和当前版本。','材料对标签：牌号、厚度、宽度和批次。','模具对编号：模具、产品和工位对应。']),
('生产部：首件与巡检','先确认，再连续做；发现趋势就处理',['首件重点：★尺寸、孔位、折弯/成型、外观、单重。','生产中重点：毛刺是否变大、尺寸是否飘、送料和卸料是否正常。','首件不合格，不进入正常生产。','频次按公司和产品实际要求确认，不在手册里硬编。']),
('生产部：异常处理四句话','现场卡片只记最有用的四件事',['先做什么：停机、复测或保留样件。','隔离多少：从最近一次确认合格之后追溯。','通知谁：生产负责人或模具部。','何时恢复：处理后重新首件或试模确认。']),
('生产部：容易混淆的材料','看标签，不看颜色；看客户和型号，不只看型号',['同型号不同客户，材料可能不同。','同客户不同零件，上机壳、下机壳、磁栅和制动片也可能不同。','0825、0832、1027、1234不能只看型号。','制动片要区分硬料、软料和料宽。']),
('模具部负责什么','让模具能生产、出问题能找到原因',['准备模具、装配和试模。','建立模具、工位、镶件和BOM资料。','处理断冲、崩刃、毛刺、送料、导正、折弯和卸料异常。','管理常用镶件和备品。']),
('级进模：料带按工位往前走','先看排样图，再看单个镶件',['OP10：导正孔、预冲孔。','OP20：小孔、端子槽。','OP30：外形预切。','OP40：压印、压筋、预成型。','OP50：折弯、翻孔。','OP60整形 → OP70落料。']),
('镶件识别：冲头类','看到外形，先判断它在做什么',['圆形小孔冲头：小孔、导正孔。','台阶式圆冲：翻孔、拉伸、局部成型。','C型开口筒式镶件：大圆弧、外形预切、避让载体。','异形方冲/扁冲：长槽、端子槽、直边切边。']),
('镶件识别：成型与辅助类','这些零件看似不冲孔，但决定模具能否稳定工作',['带凸台成型冲：预压、压筋、压印、折弯前成型。','凹模镶块：承接冲头，磨损后单独更换。','导柱导套：保证合模导向。','弹簧与等高套筒：提供卸料力、限制行程。','导正销：进入导正孔，修正送料误差。']),
('镶件记录：一件零件一条记录','不要把一套模具写成一句话',['零件ID、名称、类型。','模具编号、工位号、装配位置。','材料、热处理要求、关键尺寸。','刃口、挂台、方向特征、图纸和版本。','备品数量、存放位置和状态。']),
('SolidWorks建模最小规范','六步够用，关键是基准一致',['1 先画料带排样，定步距和送料方向。','2 建总装配和模板基准。','3 按工位建立镶件，文件名带零件ID。','4 按实际结构做装配约束。','5 检查干涉、合模和卸料板行程。','6 出工程图和BOM，和零件台账核对。']),
('模具：拆、装、试不能跳步','不是增加手续，是避免同一个问题重复发生',['拆前拍照，零件按工位摆放。','相似镶件标方向，刃口和配合面保护好。','装配核对挂台、开口、刃口和导正方向。','手动或低速确认，再进入试模。','试模记录调整前后数据和样件结果。']),
('修模报告要回答什么','“换了一个冲头”不是完整原因',['发生了什么：工位和故障现象。','为什么发生：刃口、偏载、送料、材料或装配状态。','怎么处理：维修、调整或更换什么。','是否恢复：关联试模编号和结果。','下次怎么少出：补检查点、备品或资料。']),
('行政管理部负责什么','小公司行政重点是让信息和人不掉线',['订单和资料接收、转交（具体接单职责待确认）。','入职、转岗、离职资料。','考勤、请假和加班资料。','文件、表格、通知、培训和综合协调。']),
('行政：文件和人员资料','一件事有记录，下一位接手才不会重新猜',['新版文件写清名称、版本、生效日期和发放位置。','旧版回收或标记作废。','培训记录写清谁、学了什么、依据哪个版本。','合同、工资、社保和审批按公司正式制度执行。']),
('现场速查表','把复杂内容压缩成现场能用的一页',['任务确认单：订单、产品、图纸、材料、模具、设备。','首件/巡检单：标准、实测、结果和确认人。','异常处理卡：先做什么、隔离多少、找谁、何时恢复。','镶件/试模/修机表：让模具经验留下来。']),
('未来扩展，不影响现在使用','先把三个部门跑顺，再逐步加部门',['质量管理、工程研发、采购供应链、仓储物��、业务项目、设备机修、财务。','新增部门只补职责、交接点和速查表。','现有订单主流程不重写。']),
('今天先记住三件事','新员工能说清这三件事，就已经找到工作入口',['我所在的部门负责哪一段。','我手上的产品、材料、图纸和模具是否对应。','发现问题时，我先控制现场，再找对的人。'])]
for i,(t,sub,items) in enumerate(slides,2):
 s=slide_title(t,sub,i)
 if i==4:
  # process
  steps=['客户订单','生产确认','材料准备','模具准备','首件确认','批量生产','异常处理','完工交付']
  x=0.8
  for j,st in enumerate(steps):
   rect(s,x,2.35,1.35,1.05,BLUE if j%2==0 else TEAL,BLUE,True); text(s,st,x+0.08,2.63,1.19,0.3,13,WHITE,True,align=PP_ALIGN.CENTER)
   if j<len(steps)-1: text(s,'→',x+1.39,2.63,0.35,0.3,18,ORANGE,True,align=PP_ALIGN.CENTER)
   x+=1.53
  text(s,'收到客户订单→行政管理部｜生产确认→生产部+模具部｜准备材料→生产部｜准备模具→模具部｜首件确认→生产部+模具部｜批量生产→生产部｜异常处理→生产部+模具部｜完工、记录和交付→生产部+行政管理部',0.9,4.25,11.55,0.95,13,INK,True,align=PP_ALIGN.CENTER)
  text(s,'具体接单人、交付人和确认权限以鑫犇实际安排为准。',1.2,5.45,10.9,0.35,11,MUTED,False,align=PP_ALIGN.CENTER)
 elif i in [3,10,12,21,22]:
  for j,it in enumerate(items):
   y=2.05+j*0.82; rect(s,0.85,y,0.48,0.48,ORANGE,ORANGE,True); text(s,str(j+1),0.85,y+0.08,0.48,0.25,15,NAVY,True,align=PP_ALIGN.CENTER); text(s,it,1.55,y-0.02,10.7,0.55,18,INK,j==0)
 else:
  rect(s,0.78,1.95,11.75,4.35,PALE,PALE,True); bullets(s,items,1.15,2.35,10.95,3.5,20)
prs.save(OUT)
print(OUT)
print(XLS)
print('slides',len(prs.slides),'sheets',len(sheets))
