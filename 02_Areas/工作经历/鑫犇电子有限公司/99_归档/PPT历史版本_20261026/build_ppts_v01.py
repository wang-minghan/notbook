from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path
import json, shutil
ROOT=Path(r'C:/Users/32027/OneDrive/notebook/02_Areas/工作经历/鑫犇电子有限公司/06_公司公共')

def rgb(h): return RGBColor.from_string(h.replace('#',''))
def box(sl,x,y,w,h,text,size=18,color='1F2D3D',bold=False,fill=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)) if fill else sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill); sh.line.color.rgb=rgb(fill)
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    p=tf.paragraphs[0]; p.text=text; p.font.name='Microsoft YaHei'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color); p.alignment=PP_ALIGN.CENTER
    return sh
def base(prs,bg='F7F9FC'):
    sl=prs.slides.add_slide(prs.slide_layouts[6]); sh=sl.background; sh.fill.solid(); sh.fill.fore_color.rgb=rgb(bg); return sl
def footer(sl,text,dark=False): box(sl,.55,7.12,12.2,.2,text,8,'B5C2D0' if dark else '8A97A8')
def title(sl,k,t,sub='',dark=False):
    box(sl,.6,.35,12,.3,k,10,'36C2C9' if dark else '1C9C8C',True)
    box(sl,.6,.7,12,.55,t,27,'FFFFFF' if dark else '1F2D3D',True)
    if sub: box(sl,.6,1.3,12,.35,sub,12,'C8D4DF' if dark else '526173')
def cards(sl,items, dark=False):
    n=len(items); gap=.22; w=(12.0-gap*(n-1))/n
    for i,(head,body) in enumerate(items):
        x=.65+i*(w+gap); box(sl,x,2.05,w,3.7,head+'\n\n'+body,16,'FFFFFF' if dark else '1F2D3D',True,'16324A' if dark else 'FFFFFF')

def company():
    out=ROOT/'公司介绍PPT'/'build'; out.mkdir(exist_ok=True)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    hero=ROOT/'公司介绍PPT'/'visual_assets'/'company_hero_v01.jpg'
    sl=base(prs,'111820'); sl.shapes.add_picture(str(hero), Inches(6.5),0, width=Inches(6.83),height=Inches(7.5)); title(sl,'XINBEN ELECTRONICS','鑫犇电子有限公司','精密金属零件与冲压制造能力介绍',True); box(sl,.65,2.15,5.4,1.25,'面向产品、工程、制造与品质协同的\n制造能力底稿',24,'FFFFFF',True); box(sl,.65,6.65,5.5,.35,'AI 生成概念视觉，不代表真实工厂或设备',10,'F6C453'); footer(sl,'对外介绍草案｜正式公司名称与定位待确认',True)
    data=[('02 / OVERVIEW','我们是谁','以产品、工程、制造、品质四条线理解公司能力。\n正式对外定位：待确认'),('03 / PRODUCTS','产品与应用','机壳、磁性/结构件、制动片及配套零件。\n公开产品范围：待确认'),('04 / MATERIALS','材料与版本管理','客户版本 → 图纸 → 材料 → 单重 → 检验\n敏感客户信息不在本版展示'),('05 / ENGINEERING','模具与工程技术','模具、工程技术、生产与品质共同支撑交付。\n量产能力证据：待补'),('06 / FLOW','订单到交付','需求评审 → 工程准备 → 生产 → 质量 → 交付\n责任与周期：待确认'),('07 / QUALITY','质量控制','图纸、材料、关键尺寸、外观、表面处理、单重与异常闭环。\n指标：待确认'),('08 / TRACEABILITY','产品一致性与追溯','图纸版本、材料批次、产品编码、检验记录\n真实脱敏样例：待补'),('09 / FACTORY','现场能力','工厂、设备、模具、检验现场\n待补实拍｜不以AI图替代事实照片'),('10 / COOPERATION','合作方式','需求评审 → 试制 → 量产 → 交付\n合作周期与承诺：待确认'),('11 / ROADMAP','能力扩展方向','仅展示已确认的建设事项。\n现状/建设中/拟建立：待确认'),('12 / CONTACT','合作入口','正式公司全称、地址、电话、邮箱、二维码\n全部待确认')]
    for i,(k,t,b) in enumerate(data):
        dark=i%3==0; sl=base(prs,'111820' if dark else 'F7F9FC'); title(sl,k,t,'单页唯一结论与证据边界',dark)
        cards(sl,[('核心判断',b),('资料边界','资料依据：OneDrive 公司资料\n未确认信息保持待确认'),('下一步','补齐公开事实、实拍素材与联系方式')],dark); footer(sl,'资料依据：公司事实底稿、产品信息总览、组织与职责矩阵｜草案',dark)
    prs.save(out/'xinben_company_intro_v01.pptx'); print('company',len(prs.slides),out/'xinben_company_intro_v01.pptx')

def governance():
    proj=ROOT/'公司治理PPT'/'xinben_governance_20261007'; out=proj/'build_v4'; out.mkdir(exist_ok=True)
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    baseline=json.loads((proj/'build_v4'/'baseline_texts.json').read_text(encoding='utf-8'))
    titles=[]
    for item in baseline:
        if isinstance(item,dict): titles.append(item)
    special={6:('组织架构与决策边界',['惠州市鑫犇电子有限公司','董事长｜总经理','生产部','品质部','营运部','行政部','财务部']),8:('部门职责地图',['生产部：工程/冲压/清洗/模具维修','品质部：制程巡检/成品出货检验','营运部：业务/采购/仓储','行政部：人事/后勤','财务部：会计/出纳']),11:('授权审批矩阵',['订单评审｜A：营运部｜C：专业部门（事前咨询）','采购与付款｜R：采购/财务｜A：按正式授权文件','文件发布与异常｜经办、咨询、批准分列']),32:('部门协同 RACI 示例',['采购验收：R 品质部；仓储收货协同','付款：R 财务部','每行一个 A；C=事前咨询；I=知会'])}
    for i in range(1,37):
        sl=base(prs); dark=i in (1,2,3,4,5,6)
        if i in special:
            t,items=special[i]; title(sl,f'{i:02d} / GOVERNANCE',t,'正式组织与职责依据 OneDrive 受控资料')
            for j,it in enumerate(items): box(sl,.75,2.0+(j%5)*.75,11.8,.55,it,17,'FFFFFF' if dark else '1F2D3D',True,'16324A' if dark else 'FFFFFF')
        else:
            rec=titles[i-1] if i-1<len(titles) else {}
            t=rec.get('title',rec.get('name',f'治理页面 {i}')) if isinstance(rec,dict) else f'治理页面 {i}'
            title(sl,f'{i:02d} / GOVERNANCE',str(t),'治理框架、部门协同与内控说明')
            cards(sl,[('核心机制','责任、流程、记录与复盘形成闭环。'),('事实边界','未正式发布内容标注：待确认/拟建立。'),('落地提示','具体执行以正式制度、授权文件和表单为准。')],False)
        footer(sl,'鑫犇电子公司治理与部门协同管理体系｜修正版草稿｜以正式授权文件为准')
    p=out/'xinben_governance_V4_draft.pptx'; prs.save(p); print('governance',len(prs.slides),p)
company(); governance()
