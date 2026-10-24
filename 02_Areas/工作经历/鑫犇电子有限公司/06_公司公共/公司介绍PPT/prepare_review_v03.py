from pathlib import Path
import json, hashlib, re
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw
import win32com.client
ROOT=Path(__file__).resolve().parent
SRC=ROOT/'build/xinben_company_intro_v02.pptx'
OUT=ROOT/'build/xinben_company_intro_v03_review.pptx'
QA=ROOT/'build/qa_v03_review'
QA.mkdir(exist_ok=True)
p=Presentation(str(SRC))
# Fix slide 6 handoffs to reach the next node, not empty space.
flow=[s for s in p.slides[5].shapes if s.shape_type==MSO_SHAPE_TYPE.LINE]
assert len(flow)==4, len(flow)
for i,s in enumerate(flow):
 s.begin_x=Inches(4.1+i*2.05);s.begin_y=Inches(2.08+i*.82)
 s.end_x=Inches(4.6+i*2.05);s.end_y=Inches(2.08+(i+1)*.82)
# Native editable return path from reinspection to release decision.
sl=p.slides[6]
for a,b in [((7.45,5.62),(7.65,5.62)),((7.65,5.62),(7.65,5.02)),((7.65,5.02),(7.45,5.02))]:
 s=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]))
 s.line.color.rgb=RGBColor.from_string('7967B8');s.line.width=Pt(2)
connectors=0
for i,sl in enumerate(p.slides,1):
 for s in sl.shapes:
  if s.shape_type==MSO_SHAPE_TYPE.LINE:
   ln=s._element.spPr.get_or_add_ln()
   for old in list(ln):
    if old.tag.endswith('}tailEnd'): ln.remove(old)
   tip=OxmlElement('a:tailEnd');tip.set('type','triangle');tip.set('w','med');tip.set('len','med');ln.append(tip);connectors+=1
 # Persistent review status; do not imply external publication approval.
 box=sl.shapes.add_textbox(Inches(9.35),Inches(.27),Inches(3.35),Inches(.27))
 tf=box.text_frame;tf.margin_top=0;tf.margin_bottom=0
 r=tf.paragraphs[0].add_run();r.text='内部审阅草案 · 事实未冻结';r.font.name='Microsoft YaHei';r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('B07813' if i not in (1,12) else 'F6C453')
 sl.notes_slide.notes_text_frame.text=f'第{i:02d}页｜内部审阅草案，不得作为已批准宣传资料。\n依据：company_intro_fact_base_v01.md、page_contracts_v01.md。\n未确认事项以页面待确认/待补为准；组织线索不等于产能证明。AI封面仅为概念视觉；现场照片不得用AI替代。'
p.save(str(OUT))
q=Presentation(str(OUT))
assert len(q.slides)==12
texts=['\n'.join(s.text for s in sl.shapes if s.has_text_frame) for sl in q.slides]
assert all(t.strip() and '\ufffd' not in t for t in texts)
assert all('内部审阅草案' in t for t in texts)
assert all(sl.has_notes_slide for sl in q.slides)
with __import__('zipfile').ZipFile(OUT) as z:
 arrow_count=sum(z.read(n).count(b'type="triangle"') for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n))
assert arrow_count==connectors,(arrow_count,connectors)
app=win32com.client.Dispatch('PowerPoint.Application')
deck=None
try:
 deck=app.Presentations.Open(str(OUT),ReadOnly=True,WithWindow=False)
 assert deck.Slides.Count==12
 for i in range(1,13): deck.Slides(i).Export(str(QA/f'slide_{i:02d}.jpg'),'JPG',1600,900)
finally:
 if deck: deck.Close()
 app.Quit()
imgs=[QA/f'slide_{i:02d}.jpg' for i in range(1,13)]
assert all(f.is_file() for f in imgs)
canvas=Image.new('RGB',(1920,920),'#E8E8E8');d=ImageDraw.Draw(canvas)
for j,f in enumerate(imgs):
 im=Image.open(f);im.thumbnail((470,265));x=(j%4)*480;y=(j//4)*305
 canvas.paste(im,(x+5,y+30));d.text((x+8,y+7),f'{j+1:02d}',fill='black')
canvas.save(QA/'contact_sheet_numeric.jpg',quality=95)
report={'source':str(SRC),'output':str(OUT),'slides':len(q.slides),'notes':sum(s.has_notes_slide for s in q.slides),'native_direction_markers':arrow_count,'rendered':len(imgs),'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'visual_review':'PENDING','release_status':'INTERNAL_REVIEW_ONLY','flow_handoffs_fixed':4}
(QA/'verification.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
