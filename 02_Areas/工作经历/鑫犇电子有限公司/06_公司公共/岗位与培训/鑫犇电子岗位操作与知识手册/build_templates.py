from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = Path(r"C:\Users\32027\.hermes\workspace\鑫犇电子岗位操作与知识手册\鑫犇电子_流程规范配套表格_第一版.xlsx")

SHEETS = {
    "01产品主数据": ["产品编码","客户名称","产品名称","现场常用名","产品型号","图号","图纸版本","材料牌号","材料厚度(mm)","材料宽度(mm)","单重参考(g)","模具编号","状态","生效日期","备注"],
    "02易混材料": ["对照编号","产品编码","客户名称","产品名称","产品型号","正确材料牌号","材料厚度(mm)","材料宽度(mm)","易混对象","主要区别","核对依据","混用后果","状态","备注"],
    "03任务确认": ["确认单号","生产日期","班次","工单号","产品编码","客户名称","产品名称","产品型号","图号","图纸版本","计划数量","交期","材料牌号","材料厚度(mm)","材料宽度(mm)","模具编号","设备编号","上次异常","三项一致确认","确认人","备注"],
    "04领料核对": ["记录编号","日期","工单号","产品编码","材料牌号","材料厚度(mm)","材料宽度(mm)","材料批次","领料重量(kg)","余料重量(kg)","标签状态","核对结果","核对人","备注"],
    "05首件检查": ["记录编号","日期","工单号","产品编码","图号","图纸版本","材料牌号","检查项目","图纸标准","实测值","单位","检具","结果","检查人","确认人","备注"],
    "06生产巡检": ["记录编号","日期","工单号","产品编码","检查时间","累计数量","检查项目","图纸标准","实测值","单位","外观结果","异常情况","处理结果","检查人","备注"],
    "07异常处理": ["异常编号","发现日期时间","工单号","产品编码","模具编号","工位号","异常类别","异常现象","发现时数量","隔离范围","临时处理","原因分析","处理措施","恢复条件","处理结果","记录人","确认人","关闭日期","状态"],
    "08模具零件": ["零件ID","模具编号","工位号","零件名称","零件类型","材料","热处理要求","总长(mm)","最大外径或长(mm)","最大宽度(mm)","刃口尺寸","刃口高度(mm)","挂台尺寸","特征","配合间隙","装配位置","图纸号","图纸版本","备品数量","存放位置","状态","备注"],
    "09试模记录": ["试模编号","试模日期","模具编号","产品编码","材料牌号","材料厚度(mm)","设备编号","工位号","问题位置","问题现象","调整内容","调整前数据","调整后数据","单位","样件结果","后续处理","记录人","确认人","状态","备注"],
    "10修机报告": ["报告编号","报修日期","模具编号","产品编码","工位号","故障现象","停机时数量","隔离范围","测量数据","原因分析","处理方式","更换零件ID","维修前照片路径","维修后照片路径","试模编号","试模结果","后续提升","维修人","确认人","关闭日期","状态"],
    "11模具备品": ["库存记录号","零件ID","模具编号","工位号","零件名称","最低库存","当前库存","单位","存放位置","最近入库日期","最近领用日期","补充状态","负责人","备注"],
    "12文件版本": ["文件编号","文件名称","关联对象编码","文件类型","版本号","生效日期","变更内容","发放位置","旧版处理","确认人","状态","备注"],
    "13培训记录": ["培训编号","培训日期","培训主题","关联流程编号","部门","员工编号","员工姓名","培训方式","考核结果","岗位确认结果","培训人","确认人","备注"],
    "14固定资产": ["资产编号","资产名称","资产类别","规格型号","使用部门","存放位置","责任人","启用日期","保养要求","当前状态","最近盘点日期","备注"],
}

# 所有示例均为虚拟数据，只用于说明填写方法，禁止直接作为正式参数使用。
EXAMPLES = {
    "08模具零件": [
        ["EX-P-001","EX-M-001","OP10","Φ2.0导正孔冲头","冲头","SKH51","参考：按批准图纸","80","6","","Φ2.000（示例）","6","Φ6×3（示例）","尾部挂台","待批准图纸确认","上模固定板","EX-DWG-P001","A","2","示例柜-A1","示例/待确认","虚拟示例，不可用于加工"],
        ["EX-D-001","EX-M-001","OP30","C型外圆冲头","冲头","DC53","参考：按批准图纸","65","25","","C型轮廓（见图）","5","Φ25×4（示例）","轴向开口槽","待试模确认","上模固定板","EX-DWG-D001","A","1","示例柜-A2","示例/待确认","虚拟示例，不可用于加工"],
    ],
    "09试模记录": [
        ["EX-TR-001","2026-09-18","EX-M-001","EX-LDW-1027-上机壳","C2800","0.23","EX-PR-01","OP10","导正孔","孔位向送料方向偏移","调整送料步距并复查导正销","偏移0.08（示例）","偏移0.01（示例）","mm","示例：合格","连续样件复核","张三（示例）","李四（示例）","示例/已关闭","虚拟示例"],
        ["EX-TR-002","2026-09-18","EX-M-001","EX-LDW-1027-上机壳","C2800","0.23","EX-PR-01","OP50","折弯位置","折弯高度偏高","调整成型镶件垫片","高0.06（示例）","高0.01（示例）","mm","示例：合格","记录垫片组合","张三（示例）","李四（示例）","示例/已关闭","虚拟示例"],
    ],
    "10修机报告": [
        ["EX-RP-001","2026-09-18","EX-M-001","EX-LDW-1027-上机壳","OP20","小孔毛刺突然增大","12000","自上次合格巡检后全部待判","刃口局部崩角0.05（示例）","冲头刃口崩损（示例）","更换备用冲头并重新试模","EX-P-002","C:/示例/维修前.jpg","C:/示例/维修后.jpg","EX-TR-003","示例：连续样件合格","检查该工位冲头寿命记录","王五（示例）","李四（示例）","2026-09-18","示例/已关闭"],
        ["EX-RP-002","2026-09-18","EX-M-001","EX-LDW-1027-上机壳","OP60","整形后尺寸波动","18500","最近一次合格巡检后产品待判","尺寸波动0.07（示例）","成型镶件松动（示例）","重新定位紧固并复核关联工位","","C:/示例/维修前2.jpg","C:/示例/维修后2.jpg","EX-TR-004","示例：尺寸恢复稳定","补充紧固检查点","王五（示例）","李四（示例）","2026-09-18","示例/已关闭"],
    ],
}

README_ROWS = [
    ["文件用途", "本工作簿与《鑫犇电子岗位操作与知识手册》配套使用，PPT中的表格字段与本文件一致。"],
    ["使用对象", "生产、模具、行政管理相关员工。默认使用者熟悉现场工作，本表用于统一记录与追溯。"],
    ["数据状态", "正式数据须经公司确认；带EX-或“示例”的内容均为虚拟填写示例，不得直接用于生产、加工或检验。"],
    ["日期格式", "统一使用YYYY-MM-DD；需要时间时使用YYYY-MM-DD HH:MM。"],
    ["单位规则", "厚度/宽度/长度默认mm，重量字段按表头标明g或kg；其他单位必须在单位列填写。"],
    ["编号规则", "产品、模具、零件、异常、试模、修机等编号不得重复；如公司已有编号体系，直接沿用。"],
    ["待确认", "频次、参数、责任人、审批层级等尚未确认的内容不得自行补成正式标准。"],
]

navy = "17324D"
blue = "2F75B5"
light_blue = "D9EAF7"
orange = "F4B183"
gray = "E7E6E6"
white = "FFFFFF"
thin = Side(style="thin", color="A6A6A6")

wb = Workbook()
ws = wb.active
ws.title = "使用说明"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 105
ws["A1"] = "鑫犇电子流程规范配套表格｜第一版"
ws["A1"].font = Font(size=18, bold=True, color=white)
ws["A1"].fill = PatternFill("solid", fgColor=navy)
ws.merge_cells("A1:B1")
ws.row_dimensions[1].height = 30
for r, row in enumerate(README_ROWS, 3):
    ws.cell(r, 1, row[0]); ws.cell(r, 2, row[1])
    ws.cell(r, 1).font = Font(bold=True, color=navy)
    ws.cell(r, 1).fill = PatternFill("solid", fgColor=light_blue)
    for c in range(1, 3):
        ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(r, c).border = Border(bottom=thin)
    ws.row_dimensions[r].height = 34

for idx, (name, headers) in enumerate(SHEETS.items(), 1):
    sh = wb.create_sheet(name)
    sh.sheet_view.showGridLines = False
    sh.freeze_panes = "A2"
    sh.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for c, h in enumerate(headers, 1):
        cell = sh.cell(1, c, h)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sh.row_dimensions[1].height = 36
    for c, h in enumerate(headers, 1):
        width = max(12, min(24, len(h) * 2 + 4))
        if any(x in h for x in ["现象","原因","处理","备注","内容","范围","要求"]): width = 24
        sh.column_dimensions[get_column_letter(c)].width = width
    for row in EXAMPLES.get(name, []):
        sh.append(row)
    # Reserve usable blank rows and apply consistent formatting
    for r in range(2, 202):
        for c in range(1, len(headers)+1):
            cell = sh.cell(r, c)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9D9D9"))
        if r <= 1 + len(EXAMPLES.get(name, [])):
            for c in range(1, len(headers)+1):
                sh.cell(r,c).fill = PatternFill("solid", fgColor="FFF2CC")
    # native table spanning header + enough blank rows for immediate use
    ref = f"A1:{get_column_letter(len(headers))}201"
    tab = Table(displayName=f"T{idx:02d}", ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
    sh.add_table(tab)

# Dropdowns for recurring fields
for sheet_name in ["01产品主数据","02易混材料","07异常处理","09试模记录","10修机报告"]:
    sh = wb[sheet_name]
    status_col = None
    for cell in sh[1]:
        if cell.value == "状态": status_col = cell.column
    if status_col:
        dv = DataValidation(type="list", formula1='"在用,试制,停用,待确认,处理中,已关闭,示例/待确认,示例/已关闭"', allow_blank=True)
        sh.add_data_validation(dv)
        dv.add(f"{get_column_letter(status_col)}2:{get_column_letter(status_col)}201")

for sheet_name in ["05首件检查","06生产巡检"]:
    sh = wb[sheet_name]
    result_col = next(c.column for c in sh[1] if c.value == "结果") if sheet_name == "05首件检查" else None
    if result_col:
        dv = DataValidation(type="list", formula1='"合格,不合格,待判"', allow_blank=True)
        sh.add_data_validation(dv); dv.add(f"{get_column_letter(result_col)}2:{get_column_letter(result_col)}201")

# Visual warning for any cell containing 待确认
for sh in wb.worksheets[1:]:
    rng = f"A2:{get_column_letter(sh.max_column)}201"
    sh.conditional_formatting.add(rng, FormulaRule(formula=['ISNUMBER(SEARCH("待确认",A2))'], fill=PatternFill("solid", fgColor=orange)))
    sh.sheet_properties.pageSetUpPr.fitToPage = True
    sh.page_setup.fitToWidth = 1
    sh.page_setup.fitToHeight = 0
    sh.page_layout_view = False

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)

# Verify workbook by reopening and checking exact schema and example counts.
check = load_workbook(OUT, data_only=False)
assert check.sheetnames == ["使用说明"] + list(SHEETS)
for name, headers in SHEETS.items():
    actual = [check[name].cell(1,c).value for c in range(1, len(headers)+1)]
    assert actual == headers, (name, actual)
for name in ["08模具零件","09试模记录","10修机报告"]:
    assert check[name].cell(2,1).value and check[name].cell(3,1).value
print(f"CREATED={OUT}")
print(f"SHEETS={len(check.sheetnames)}")
print("VERIFIED=schemas_and_examples_ok")
